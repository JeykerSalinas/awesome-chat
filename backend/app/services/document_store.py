from __future__ import annotations

import json
import shutil
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Literal
from uuid import uuid4

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from app.core.config import BASE_DIR
from app.schemas.cv_assistant import ParagraphRevisionInput
from app.schemas.documents import DocumentParagraphBlock, DocumentSummary

DocumentFormat = Literal["docx", "pdf"]

DOCUMENTS_ROOT = BASE_DIR / "data" / "documents"
INDEX_PATH = DOCUMENTS_ROOT / "index.json"


def _ensure_storage() -> None:
    DOCUMENTS_ROOT.mkdir(parents=True, exist_ok=True)
    if not INDEX_PATH.exists():
        INDEX_PATH.write_text("{}", encoding="utf-8")


def _load_index() -> dict[str, dict]:
    _ensure_storage()
    return json.loads(INDEX_PATH.read_text(encoding="utf-8"))


def _save_index(index: dict[str, dict]) -> None:
    _ensure_storage()
    INDEX_PATH.write_text(json.dumps(index, ensure_ascii=True, indent=2), encoding="utf-8")


def _now() -> datetime:
    return datetime.now(UTC)


def _document_dir(document_id: str) -> Path:
    return DOCUMENTS_ROOT / document_id


def _docx_path(document_id: str) -> Path:
    return _document_dir(document_id) / "document.docx"


def _pdf_path(document_id: str) -> Path:
    return _document_dir(document_id) / "document.pdf"


def _to_summary(metadata: dict) -> DocumentSummary:
    return DocumentSummary(
        id=metadata["id"],
        display_name=metadata["display_name"],
        original_filename=metadata["original_filename"],
        kind=metadata["kind"],
        source_document_id=metadata.get("source_document_id"),
        created_at=datetime.fromisoformat(metadata["created_at"]),
        updated_at=datetime.fromisoformat(metadata["updated_at"]),
        available_formats=metadata["available_formats"],
    )


def list_documents() -> list[DocumentSummary]:
    index = _load_index()
    documents = [_to_summary(item) for item in index.values()]
    return sorted(documents, key=lambda item: item.updated_at, reverse=True)


def get_document(document_id: str) -> DocumentSummary:
    index = _load_index()
    metadata = index.get(document_id)
    if metadata is None:
        raise FileNotFoundError(f"Document {document_id} was not found.")
    return _to_summary(metadata)


def _store_metadata(metadata: dict) -> DocumentSummary:
    index = _load_index()
    index[metadata["id"]] = metadata
    _save_index(index)
    return _to_summary(metadata)


def _export_docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    document = DocxDocument(docx_path)
    styles = getSampleStyleSheet()
    body_style = styles["BodyText"]
    heading_style = ParagraphStyle(
        "CvHeading",
        parent=styles["Heading2"],
        spaceBefore=8,
        spaceAfter=4,
    )
    flow = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style_name = paragraph.style.name.lower() if paragraph.style and paragraph.style.name else ""
        style = heading_style if "heading" in style_name else body_style
        flow.append(Paragraph(text.replace("\n", "<br/>"), style))
        flow.append(Spacer(1, 0.18 * cm))

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        leftMargin=1.8 * cm,
        rightMargin=1.8 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
    )
    pdf.build(flow)


def _register_document(
    *,
    document_id: str,
    display_name: str,
    original_filename: str,
    kind: Literal["uploaded", "generated"],
    source_document_id: str | None = None,
) -> DocumentSummary:
    timestamp = _now().isoformat()
    metadata = {
        "id": document_id,
        "display_name": display_name,
        "original_filename": original_filename,
        "kind": kind,
        "source_document_id": source_document_id,
        "created_at": timestamp,
        "updated_at": timestamp,
        "available_formats": ["docx", "pdf"],
    }
    return _store_metadata(metadata)


def create_document_from_upload(
    *,
    filename: str,
    content: bytes,
    display_name: str | None = None,
) -> DocumentSummary:
    if not filename.lower().endswith(".docx"):
        raise ValueError("Only .docx files are supported for editing.")

    document_id = str(uuid4())
    target_dir = _document_dir(document_id)
    target_dir.mkdir(parents=True, exist_ok=True)

    docx_path = _docx_path(document_id)
    docx_path.write_bytes(content)
    _export_docx_to_pdf(docx_path, _pdf_path(document_id))

    base_name = Path(filename).stem
    return _register_document(
        document_id=document_id,
        display_name=display_name or base_name,
        original_filename=filename,
        kind="uploaded",
    )


def rename_document(document_id: str, display_name: str) -> DocumentSummary:
    index = _load_index()
    metadata = index.get(document_id)
    if metadata is None:
        raise FileNotFoundError(f"Document {document_id} was not found.")

    metadata["display_name"] = display_name
    metadata["updated_at"] = _now().isoformat()
    index[document_id] = metadata
    _save_index(index)
    return _to_summary(metadata)


def delete_document(document_id: str) -> None:
    index = _load_index()
    metadata = index.pop(document_id, None)
    if metadata is None:
        raise FileNotFoundError(f"Document {document_id} was not found.")
    _save_index(index)
    shutil.rmtree(_document_dir(document_id), ignore_errors=True)


def get_document_download_path(document_id: str, format: DocumentFormat) -> Path:
    get_document(document_id)
    path = _docx_path(document_id) if format == "docx" else _pdf_path(document_id)
    if not path.exists():
        raise FileNotFoundError(f"Format {format} is not available for document {document_id}.")
    return path


def extract_document_blocks(document_id: str) -> list[DocumentParagraphBlock]:
    get_document(document_id)
    document = DocxDocument(_docx_path(document_id))
    blocks: list[DocumentParagraphBlock] = []

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style and paragraph.style.name else "Normal"
        blocks.append(
            DocumentParagraphBlock(
                index=index,
                style=style_name,
                text=text,
            )
        )

    return blocks


def create_generated_variant(
    *,
    source_document_id: str,
    display_name: str,
    revisions: list[ParagraphRevisionInput],
) -> DocumentSummary:
    source_summary = get_document(source_document_id)
    source_docx_path = get_document_download_path(source_document_id, "docx")

    generated_id = str(uuid4())
    target_dir = _document_dir(generated_id)
    target_dir.mkdir(parents=True, exist_ok=True)
    output_docx_path = _docx_path(generated_id)

    document = DocxDocument(source_docx_path)
    revision_map = {revision.index: revision.text for revision in revisions}
    for index, paragraph in enumerate(document.paragraphs):
        replacement = revision_map.get(index)
        if replacement is not None:
            paragraph.text = replacement

    document.save(output_docx_path)
    _export_docx_to_pdf(output_docx_path, _pdf_path(generated_id))

    return _register_document(
        document_id=generated_id,
        display_name=display_name,
        original_filename=f"{Path(source_summary.original_filename).stem}.docx",
        kind="generated",
        source_document_id=source_document_id,
    )


def get_document_bytes(document_id: str, format: DocumentFormat) -> bytes:
    return get_document_download_path(document_id, format).read_bytes()

